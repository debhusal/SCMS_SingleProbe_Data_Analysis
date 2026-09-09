"""Single-Probe legacy processing. No pandas or all-workbook materialization."""
from __future__ import annotations

import csv
import io
import json
import math
import re
from bisect import bisect_left
from contextlib import ExitStack
from dataclasses import asdict, dataclass
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook


class SpectrumError(ValueError):
    """Input or scientific validation failure safe to show to the user."""


@dataclass(frozen=True)
class Settings:
    background_cutoff: float = 3000.0
    sample_cutoff: float = 3000.0
    tolerance_ppm: float = 5.0
    normalization_factor: float = 100000.0
    ion_mode: str = "POS"

    def validate(self):
        for key in ("background_cutoff", "sample_cutoff", "tolerance_ppm", "normalization_factor"):
            value = getattr(self, key)
            if not math.isfinite(value) or value < 0:
                raise SpectrumError(f"{key.replace('_', ' ')} must be finite and nonnegative.")
        if self.normalization_factor == 0 or not 0 < self.tolerance_ppm < 1000000:
            raise SpectrumError("Normalization must be positive; tolerance must be between 0 and 1000000 ppm.")
        if self.ion_mode not in ("POS", "NEG"):
            raise SpectrumError("Choose POS or NEG for this run.")


def header_indices(row):
    names = [re.sub(r"[^a-z0-9]", "", str(v).lower()) if v is not None else "" for v in row]
    masses = [i for i, v in enumerate(names) if v in ("mz", "masscharge", "masstocharge")]
    intensities = [i for i, v in enumerate(names) if v == "intensity"]
    if masses and intensities:
        if len(masses) != 1 or len(intensities) != 1:
            raise SpectrumError("Ambiguous header: more than one m/z or Intensity column.")
        return masses[0], intensities[0]
    return None


def peaks(rows, label):
    """Skip preamble/blank rows; reject corrupt data rather than silently drop peaks."""
    indices = None
    any_content = False
    count = 0
    for line, row in enumerate(rows, 1):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        any_content = True
        if indices is None:
            indices = header_indices(row)
            continue
        if header_indices(row) is not None:
            continue  # repeated exported header
        try:
            mass, intensity = (row[i] for i in indices)
            if isinstance(mass, bool) or isinstance(intensity, bool):
                raise ValueError
            mass, intensity = float(mass), float(intensity)
            if not math.isfinite(mass) or not math.isfinite(intensity) or mass <= 0 or intensity < 0:
                raise ValueError
        except (IndexError, TypeError, ValueError):
            raise SpectrumError(
                f"{label}, row {line}: invalid or missing numeric m/z or Intensity. "
                "Use positive m/z and finite, nonnegative intensity; remove footer text. "
                "For formulas, save calculated values in Excel first."
            ) from None
        count += 1
        yield mass, intensity
    if not any_content:
        raise SpectrumError(f"{label}: empty sheet/file.")
    if indices is None:
        raise SpectrumError(f"{label}: missing m.z and/or Intensity header. Relative is not Intensity.")
    if not count:
        raise SpectrumError(f"{label}: header found, but no spectrum peaks.")


def workbook(source):
    try:
        if hasattr(source, "seek"):
            source.seek(0)
        return load_workbook(source, read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise SpectrumError("Cannot open Excel workbook. Use an unencrypted .xlsx or .xlsm file; "
                            "open older .xls files in Excel and Save As .xlsx.") from exc


def sheet_names(source):
    book = workbook(source)
    try:
        return book.sheetnames
    finally:
        book.close()


def sheet_rows(sheet):
    # Some exporters report incorrect used dimensions. Stream the actual XML rows.
    sheet.reset_dimensions()
    return sheet.iter_rows(values_only=True)


def text_rows(source):
    source.seek(0)
    prefix = source.read(4096)
    source.seek(0)
    encoding = "utf-16" if prefix.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
    if encoding == "utf-8-sig":
        try:
            prefix.decode(encoding)
        except UnicodeDecodeError:
            encoding = "cp1252"
    text = io.TextIOWrapper(source, encoding=encoding, newline="")
    delimiter = None
    try:
        for line in text:
            if delimiter is None:
                for candidate in (",", "\t", ";", "whitespace"):
                    row = re.split(r"\s+", line.strip()) if candidate == "whitespace" else next(csv.reader([line], delimiter=candidate))
                    if header_indices(row) is not None:
                        delimiter = candidate
                        break
                else:
                    yield [line.strip()]
                    continue
            row = re.split(r"\s+", line.strip()) if delimiter == "whitespace" else next(csv.reader([line], delimiter=delimiter))
            yield row
    finally:
        text.detach()  # caller owns upload stream


def legacy_boundary(mass, ppm=5.0):
    # Use the exact literals and percent-formatting from format.py at default.
    low_factor, high_factor = (0.999995, 1.000005) if ppm == 5 else (1 - ppm / 1e6, 1 + ppm / 1e6)
    return float("%8.4f" % (mass * low_factor)), float("%8.4f" % (mass * high_factor))


class Exclusions:
    """Indexed union of OPEN intervals; touching intervals must stay separate."""
    def __init__(self, ranges):
        merged = []
        for lower, upper in sorted(ranges):
            if lower >= upper:
                continue
            if merged and lower < merged[-1][1]:
                merged[-1] = (merged[-1][0], max(upper, merged[-1][1]))
            else:
                merged.append((lower, upper))
        self.ranges = merged
        self.lower = [r[0] for r in merged]

    def contains(self, mass):
        i = bisect_left(self.lower, mass) - 1
        return i >= 0 and self.ranges[i][0] < mass < self.ranges[i][1]


def background_ranges(source, filename, settings, sheet=None):
    settings.validate()
    book = None
    rows = None
    try:
        if Path(filename).suffix.lower() in (".xlsx", ".xlsm"):
            book = workbook(source)
            if sheet is None and len(book.sheetnames) != 1:
                raise SpectrumError("Choose one background sheet; background sheets are never combined.")
            rows = sheet_rows(book[sheet or book.sheetnames[0]])
        elif Path(filename).suffix.lower() in (".csv", ".txt", ".tsv"):
            rows = text_rows(source)
        else:
            raise SpectrumError("Background must be .xlsx, .xlsm, .csv, .txt or .tsv. Convert .xls to .xlsx in Excel.")
        ranges, before = [], 0
        for mass, intensity in peaks(rows, "Background"):
            before += 1
            if intensity >= settings.background_cutoff:
                ranges.append(legacy_boundary(mass, settings.tolerance_ppm))
        return ranges, before
    finally:
        if rows is not None and hasattr(rows, "close"):
            rows.close()
        if book:
            book.close()


def r_number(value):
    # rbind with headers2 coerces legacy columns to character (R's 15 significant digits).
    return format(value, ".15g")


def process_workbook(source, ranges, settings, destination, progress=None):
    """Disk-spool one cell at a time, then transpose streams to legacy wide CSV."""
    settings.validate()
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    if any(destination.iterdir()):
        raise SpectrumError("Output directory must be empty; choose a new run directory.")
    index = Exclusions(ranges)
    summaries = []
    prefix = settings.ion_mode
    legacy = destination / f"{prefix}_legacy.csv"
    tidy = destination / f"{prefix}_tidy.csv"
    summary_path = destination / f"{prefix}_summary.csv"
    book = workbook(source)
    try:
        with TemporaryDirectory(prefix="spool_", dir=destination) as temp, tidy.open("w", newline="", encoding="utf-8") as tidy_file:
            tidy_writer = csv.writer(tidy_file)
            tidy_writer.writerow(["ion_mode", "cell_id", "sheet", "mass", "intensity", "normalized_intensity"])
            spools = []
            for cell_id, sheet in enumerate(book.worksheets, 1):
                if progress:
                    progress(cell_id - 1, len(book.sheetnames), sheet.title)
                spool = Path(temp) / f"{cell_id}.csv"
                before = removed_bg = removed_threshold = remaining = 0
                with spool.open("w", newline="", encoding="utf-8") as out:
                    writer = csv.writer(out)

                    def post_background_intensities():
                        nonlocal before, removed_bg, removed_threshold, remaining
                        for mass, intensity in peaks(sheet_rows(sheet), f"Sheet '{sheet.title}'"):
                            before += 1
                            if index.contains(mass):
                                removed_bg += 1
                                continue
                            if intensity < settings.sample_cutoff:
                                removed_threshold += 1
                            else:
                                writer.writerow([repr(mass), repr(intensity)])
                                remaining += 1
                            yield intensity  # R denominator includes below-cutoff peaks

                    total = math.fsum(post_background_intensities())
                if not math.isfinite(total) or total <= 0:
                    raise SpectrumError(f"Sheet '{sheet.title}': total intensity after background removal is zero or nonfinite; cannot normalize.")
                if remaining == 0:
                    raise SpectrumError(f"Sheet '{sheet.title}': no peaks remain after the intensity cutoff; cannot export this cell.")
                normalized_spool = Path(temp) / f"{cell_id}_normalized.csv"
                with spool.open(newline="", encoding="utf-8") as inp, normalized_spool.open("w", newline="", encoding="utf-8") as out:
                    writer = csv.writer(out)
                    for mass_text, intensity_text in csv.reader(inp):
                        mass, intensity = float(mass_text), float(intensity_text)
                        abundance = intensity / total * settings.normalization_factor
                        writer.writerow([r_number(mass), r_number(abundance)])
                        tidy_writer.writerow([prefix, cell_id, sheet.title, repr(mass), repr(intensity), repr(abundance)])
                spool.unlink()
                spools.append(normalized_spool)
                summaries.append({"cell_id": cell_id, "sheet": sheet.title, "peaks_before": before,
                                  "removed_background": removed_bg, "removed_intensity": removed_threshold,
                                  "remaining": remaining, "normalization_denominator": total})
            with ExitStack() as stack, legacy.open("w", newline="", encoding="utf-8") as out:
                writer = csv.writer(out, quoting=csv.QUOTE_ALL)
                writer.writerow([str(i) for i in range(1, len(spools) + 1) for _ in range(2)])
                writer.writerow([name for _ in spools for name in ("mass", "abund")])
                readers = [csv.reader(stack.enter_context(p.open(newline="", encoding="utf-8"))) for p in spools]
                for _ in range(max(s["remaining"] for s in summaries)):
                    row = []
                    for reader in readers:
                        pair = next(reader, None)
                        # write.table quotes character fields, but writes missing NA unquoted.
                        row.extend(['"' + value + '"' for value in pair] if pair else ["NA", "NA"])
                    out.write(",".join(row) + "\n")
        with summary_path.open("w", newline="", encoding="utf-8") as out:
            writer = csv.DictWriter(out, fieldnames=list(summaries[0]))
            writer.writeheader()
            writer.writerows(summaries)
        with (destination / f"{prefix}_background_ranges.csv").open("w", newline="", encoding="utf-8") as out:
            writer = csv.writer(out)
            writer.writerow(["lower_exclusive", "upper_exclusive"])
            writer.writerows((f"{low:.4f}", f"{high:.4f}") for low, high in ranges)
        (destination / f"{prefix}_settings.json").write_text(json.dumps({
            **asdict(settings), "normalization": "after background removal, BEFORE sample cutoff (reference R)",
            "boundaries": "%8.4f; strict lower < mass < upper", "sheets": book.sheetnames,
            "legacy_padding": "NA", "legacy_numeric_serialization": "15 significant digits",
        }, indent=2), encoding="utf-8")
        if progress:
            progress(len(summaries), len(summaries), "Complete")
        return summaries
    except Exception:
        # Never leave apparently complete partial exports after a validation failure.
        for path in destination.glob(f"{prefix}_*"):
            if path.is_file():
                path.unlink()
        raise
    finally:
        book.close()
